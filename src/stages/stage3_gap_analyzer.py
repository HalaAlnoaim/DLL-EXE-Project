"""
Stage 3: Gap Analyzer
Analyzes detection coverage gaps against the MITRE ATT&CK framework.
"""

import json
import os
import pandas as pd
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Set
import sys

# Add utils to path for MITRE loader
sys.path.append(os.path.join(os.path.dirname(__file__), '..', 'utils'))
from mitre_loader import MitreAttackLoader


class GapAnalyzer:
    """Analyzes detection coverage gaps against MITRE ATT&CK framework."""
    
    def __init__(self, base_output_dir, logger):
        """
        Initialize gap analyzer.
        
        Args:
            base_output_dir (str): Base output directory
            logger: Logger instance
        """
        self.base_output_dir = base_output_dir
        self.logger = logger
        self.output_folder = os.path.join(base_output_dir, "03_gap_analysis")
        
        # Initialize MITRE ATT&CK loader
        mitre_cache_dir = os.path.join(base_output_dir, "data", "mitre")
        self.mitre_loader = MitreAttackLoader(mitre_cache_dir)
        
    def load_client_excel(self, excel_file_path: str) -> pd.DataFrame:
        """
        Load client detection rules from Excel file.
        
        Args:
            excel_file_path (str): Path to Excel file
            
        Returns:
            pd.DataFrame: Client rules data
        """
        try:
            # Load the main detection rules sheet
            df = pd.read_excel(excel_file_path, sheet_name='Detection Rules')
            self.logger.info(f"Loaded {len(df)} rules from {excel_file_path}")
            return df
        except Exception as e:
            self.logger.error(f"Failed to load Excel file {excel_file_path}: {str(e)}")
            return pd.DataFrame()
            
    def extract_covered_techniques(self, df: pd.DataFrame) -> Set[str]:
        """
        Extract MITRE techniques covered by client rules.
        
        Args:
            df (pd.DataFrame): Client rules dataframe
            
        Returns:
            set: Set of covered technique IDs
        """
        covered_techniques = set()
        
        # Extract main techniques
        main_techniques = df['Technique ID'].dropna().unique()
        for tech_id in main_techniques:
            if tech_id and str(tech_id).strip():
                covered_techniques.add(str(tech_id).strip())
                
        # Extract sub-techniques
        sub_techniques = df['Sub-technique ID'].dropna().unique()
        for sub_tech_id in sub_techniques:
            if sub_tech_id and str(sub_tech_id).strip():
                covered_techniques.add(str(sub_tech_id).strip())
                
        self.logger.info(f"Found {len(covered_techniques)} unique covered techniques")
        return covered_techniques
        
    def analyze_client_coverage(self, client_excel_path: str) -> Dict:
        """
        Analyze coverage for a single client.
        
        Args:
            client_excel_path (str): Path to client Excel file
            
        Returns:
            dict: Coverage analysis results
        """
        client_name = Path(client_excel_path).stem
        self.logger.info(f"Analyzing coverage for {client_name}")
        
        # Load client rules
        client_df = self.load_client_excel(client_excel_path)
        if client_df.empty:
            return {}
            
        # Get covered techniques
        covered_techniques = self.extract_covered_techniques(client_df)
        
        # Load MITRE data
        mitre_data = self.mitre_loader.get_processed_data()
        if not mitre_data:
            self.logger.error("Failed to load MITRE ATT&CK data")
            return {}
            
        # Get all MITRE techniques
        all_mitre_techniques = self.mitre_loader.get_all_techniques(include_subtechniques=True)
        all_technique_ids = set(all_mitre_techniques.keys())
        
        # Calculate gaps
        uncovered_techniques = all_technique_ids - covered_techniques
        
        # Analyze by tactic
        tactic_analysis = self._analyze_by_tactic(
            covered_techniques, 
            uncovered_techniques, 
            all_mitre_techniques
        )
        
        # Generate coverage statistics
        coverage_stats = {
            "client_name": client_name,
            "analysis_date": datetime.now().isoformat(),
            "total_mitre_techniques": len(all_technique_ids),
            "covered_techniques": len(covered_techniques),
            "uncovered_techniques": len(uncovered_techniques),
            "coverage_percentage": round((len(covered_techniques) / len(all_technique_ids)) * 100, 2),
            "covered_technique_list": sorted(list(covered_techniques)),
            "uncovered_technique_list": sorted(list(uncovered_techniques)),
            "tactic_analysis": tactic_analysis
        }
        
        # Add detailed gap information
        coverage_stats["gap_details"] = self._generate_gap_details(
            uncovered_techniques, 
            all_mitre_techniques
        )
        
        return coverage_stats
        
    def _analyze_by_tactic(self, covered: Set[str], uncovered: Set[str], all_techniques: Dict) -> Dict:
        """Analyze coverage by MITRE tactic."""
        tactic_analysis = {}
        
        # Group techniques by tactic
        tactic_groups = {}
        for tech_id, tech_data in all_techniques.items():
            tactics = tech_data.get("tactics", [])
            for tactic in tactics:
                if tactic not in tactic_groups:
                    tactic_groups[tactic] = set()
                tactic_groups[tactic].add(tech_id)
                
        # Calculate coverage for each tactic
        for tactic, tactic_techniques in tactic_groups.items():
            covered_in_tactic = tactic_techniques.intersection(covered)
            uncovered_in_tactic = tactic_techniques.intersection(uncovered)
            
            coverage_pct = (len(covered_in_tactic) / len(tactic_techniques)) * 100 if tactic_techniques else 0
            
            tactic_analysis[tactic] = {
                "total_techniques": len(tactic_techniques),
                "covered_techniques": len(covered_in_tactic),
                "uncovered_techniques": len(uncovered_in_tactic),
                "coverage_percentage": round(coverage_pct, 2),
                "covered_technique_ids": sorted(list(covered_in_tactic)),
                "uncovered_technique_ids": sorted(list(uncovered_in_tactic))
            }
            
        return tactic_analysis
        
    def _generate_gap_details(self, uncovered_techniques: Set[str], all_techniques: Dict) -> List[Dict]:
        """Generate detailed information about coverage gaps."""
        gap_details = []
        
        for tech_id in sorted(uncovered_techniques):
            tech_data = all_techniques.get(tech_id, {})
            
            gap_info = {
                "technique_id": tech_id,
                "technique_name": tech_data.get("name", "Unknown"),
                "tactics": tech_data.get("tactics", []),
                "platforms": tech_data.get("platforms", []),
                "data_sources": tech_data.get("data_sources", []),
                "description": tech_data.get("description", "")[:200] + "..." if len(tech_data.get("description", "")) > 200 else tech_data.get("description", ""),
                "mitre_url": tech_data.get("url", ""),
                "is_subtechnique": "." in tech_id,
                "parent_technique": tech_id.split('.')[0] if "." in tech_id else None
            }
            
            gap_details.append(gap_info)
            
        return gap_details
        
    def create_gap_report(self, coverage_analysis: Dict, output_path: str):
        """
        Create detailed gap analysis report in Excel format.
        
        Args:
            coverage_analysis (dict): Coverage analysis results
            output_path (str): Output file path
        """
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # Summary sheet
            self._create_summary_sheet(coverage_analysis, writer)
            
            # Gap details sheet
            self._create_gap_details_sheet(coverage_analysis, writer)
            
            # Tactic analysis sheet
            self._create_tactic_analysis_sheet(coverage_analysis, writer)
            
            # Recommendations sheet
            self._create_recommendations_sheet(coverage_analysis, writer)
            
            # Format the workbook
            self._format_gap_report(writer)
            
    def _create_summary_sheet(self, analysis: Dict, writer):
        """Create summary sheet for gap analysis."""
        summary_data = [
            ["S-Rank Core - Coverage Gap Analysis"],
            [""],
            ["Client Information", ""],
            ["Client Name", analysis.get("client_name", "Unknown")],
            ["Analysis Date", analysis.get("analysis_date", "")],
            [""],
            ["Coverage Statistics", ""],
            ["Total MITRE Techniques", analysis.get("total_mitre_techniques", 0)],
            ["Covered Techniques", analysis.get("covered_techniques", 0)],
            ["Uncovered Techniques", analysis.get("uncovered_techniques", 0)],
            ["Coverage Percentage", f"{analysis.get('coverage_percentage', 0)}%"],
            [""],
            ["Priority Recommendations", ""],
            ["High Priority Gaps", self._count_high_priority_gaps(analysis)],
            ["Critical Tactics Gaps", self._count_critical_tactic_gaps(analysis)],
            [""],
            ["Coverage by Tactic", ""],
        ]
        
        # Add tactic coverage summary
        tactic_analysis = analysis.get("tactic_analysis", {})
        for tactic, tactic_data in tactic_analysis.items():
            coverage_pct = tactic_data.get("coverage_percentage", 0)
            summary_data.append([f"  {tactic}", f"{coverage_pct:.1f}%"])
            
        summary_df = pd.DataFrame(summary_data)
        summary_df.to_excel(writer, sheet_name='Summary', index=False, header=False)
        
    def _create_gap_details_sheet(self, analysis: Dict, writer):
        """Create detailed gap analysis sheet."""
        gap_details = analysis.get("gap_details", [])
        
        if gap_details:
            gap_df = pd.DataFrame(gap_details)
            # Reorder columns for better readability
            column_order = [
                "technique_id", "technique_name", "tactics", "platforms",
                "data_sources", "is_subtechnique", "parent_technique",
                "description", "mitre_url"
            ]
            gap_df = gap_df.reindex(columns=column_order)
            gap_df.to_excel(writer, sheet_name='Gap Details', index=False)
        else:
            # Create empty sheet with headers
            empty_df = pd.DataFrame(columns=[
                "technique_id", "technique_name", "tactics", "platforms",
                "data_sources", "is_subtechnique", "parent_technique",
                "description", "mitre_url"
            ])
            empty_df.to_excel(writer, sheet_name='Gap Details', index=False)
            
    def _create_tactic_analysis_sheet(self, analysis: Dict, writer):
        """Create tactic-level analysis sheet."""
        tactic_analysis = analysis.get("tactic_analysis", {})
        
        tactic_data = []
        for tactic, data in tactic_analysis.items():
            tactic_data.append({
                "MITRE Tactic": tactic,
                "Total Techniques": data.get("total_techniques", 0),
                "Covered Techniques": data.get("covered_techniques", 0),
                "Uncovered Techniques": data.get("uncovered_techniques", 0),
                "Coverage Percentage": f"{data.get('coverage_percentage', 0):.1f}%",
                "Coverage Status": self._get_coverage_status(data.get("coverage_percentage", 0))
            })
            
        tactic_df = pd.DataFrame(tactic_data)
        tactic_df = tactic_df.sort_values("Coverage Percentage")
        tactic_df.to_excel(writer, sheet_name='Tactic Analysis', index=False)
        
    def _create_recommendations_sheet(self, analysis: Dict, writer):
        """Create recommendations sheet."""
        recommendations = self._generate_recommendations(analysis)
        
        rec_df = pd.DataFrame(recommendations)
        rec_df.to_excel(writer, sheet_name='Recommendations', index=False)
        
    def _generate_recommendations(self, analysis: Dict) -> List[Dict]:
        """Generate actionable recommendations based on gap analysis."""
        recommendations = []
        
        tactic_analysis = analysis.get("tactic_analysis", {})
        gap_details = analysis.get("gap_details", [])
        
        # Find tactics with lowest coverage
        low_coverage_tactics = []
        for tactic, data in tactic_analysis.items():
            coverage_pct = data.get("coverage_percentage", 0)
            if coverage_pct < 50:  # Less than 50% coverage
                low_coverage_tactics.append((tactic, coverage_pct, data.get("uncovered_techniques", 0)))
                
        # Sort by coverage percentage (lowest first)
        low_coverage_tactics.sort(key=lambda x: x[1])
        
        # Generate recommendations for low coverage tactics
        for tactic, coverage_pct, uncovered_count in low_coverage_tactics[:5]:  # Top 5 priorities
            recommendations.append({
                "Priority": "High",
                "Category": "Tactic Coverage",
                "MITRE Tactic": tactic,
                "Current Coverage": f"{coverage_pct:.1f}%",
                "Gap Count": uncovered_count,
                "Recommendation": f"Focus on improving {tactic} detection coverage. Consider implementing rules for the {uncovered_count} uncovered techniques.",
                "Impact": "High - Addresses significant coverage gap in critical attack phase"
            })
            
        # Generate recommendations for high-impact techniques
        high_impact_techniques = ["T1003", "T1055", "T1059", "T1078", "T1190"]  # Common high-impact techniques
        
        for gap in gap_details:
            tech_id = gap.get("technique_id", "")
            if tech_id in high_impact_techniques:
                recommendations.append({
                    "Priority": "Critical",
                    "Category": "High-Impact Technique",
                    "MITRE Tactic": ", ".join(gap.get("tactics", [])),
                    "Technique": f"{tech_id} - {gap.get('technique_name', '')}",
                    "Gap Count": 1,
                    "Recommendation": f"Implement detection for {gap.get('technique_name', '')} - commonly used by attackers",
                    "Impact": "Critical - High-frequency attack technique with significant business impact"
                })
                
        # Add general recommendations
        total_gaps = analysis.get("uncovered_techniques", 0)
        if total_gaps > 100:
            recommendations.append({
                "Priority": "Medium",
                "Category": "Overall Strategy",
                "MITRE Tactic": "All",
                "Current Coverage": f"{analysis.get('coverage_percentage', 0):.1f}%",
                "Gap Count": total_gaps,
                "Recommendation": "Consider implementing a systematic approach to coverage improvement, focusing on 3-5 techniques per month",
                "Impact": "Medium - Gradual improvement in overall security posture"
            })
            
        return recommendations
        
    def _count_high_priority_gaps(self, analysis: Dict) -> int:
        """Count high priority gaps (techniques in critical tactics)."""
        critical_tactics = ["initial-access", "execution", "persistence", "credential-access"]
        
        high_priority_count = 0
        tactic_analysis = analysis.get("tactic_analysis", {})
        
        for tactic, data in tactic_analysis.items():
            if tactic.lower().replace(" ", "-") in critical_tactics:
                high_priority_count += data.get("uncovered_techniques", 0)
                
        return high_priority_count
        
    def _count_critical_tactic_gaps(self, analysis: Dict) -> int:
        """Count tactics with less than 30% coverage."""
        critical_gaps = 0
        tactic_analysis = analysis.get("tactic_analysis", {})
        
        for tactic, data in tactic_analysis.items():
            if data.get("coverage_percentage", 0) < 30:
                critical_gaps += 1
                
        return critical_gaps
        
    def _get_coverage_status(self, coverage_pct: float) -> str:
        """Get coverage status based on percentage."""
        if coverage_pct >= 80:
            return "Excellent"
        elif coverage_pct >= 60:
            return "Good"
        elif coverage_pct >= 40:
            return "Fair"
        elif coverage_pct >= 20:
            return "Poor"
        else:
            return "Critical"
            
    def _format_gap_report(self, writer):
        """Apply formatting to gap analysis report."""
        from openpyxl.styles import Font, PatternFill, Alignment
        
        # Format summary sheet
        summary_ws = writer.sheets['Summary']
        title_font = Font(bold=True, size=14)
        summary_ws['A1'].font = title_font
        
        # Format other sheets with headers
        for sheet_name in ['Gap Details', 'Tactic Analysis', 'Recommendations']:
            if sheet_name in writer.sheets:
                ws = writer.sheets[sheet_name]
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                
                # Format header row
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=col)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
    def analyze_all_clients(self, input_folder: str) -> str:
        """
        Analyze coverage gaps for all client Excel files.
        
        Args:
            input_folder (str): Folder containing client Excel files
            
        Returns:
            str: Path to output folder containing analysis reports
        """
        self.logger.info(f"Starting gap analysis from {input_folder}")
        
        # Ensure MITRE data is available
        self.logger.info("Loading MITRE ATT&CK data...")
        mitre_data = self.mitre_loader.get_processed_data()
        if not mitre_data:
            self.logger.error("Failed to load MITRE ATT&CK data")
            return self.output_folder
            
        # Find all Excel files
        excel_files = list(Path(input_folder).glob("*.xlsx"))
        if not excel_files:
            self.logger.warning(f"No Excel files found in {input_folder}")
            return self.output_folder
            
        # Ensure output directory exists
        os.makedirs(self.output_folder, exist_ok=True)
        
        successful_analyses = 0
        
        for excel_file in excel_files:
            try:
                # Analyze client coverage
                coverage_analysis = self.analyze_client_coverage(str(excel_file))
                
                if coverage_analysis:
                    # Generate output filename
                    client_name = excel_file.stem
                    timestamp = datetime.now().strftime("%Y%m%d")
                    report_filename = f"{client_name}_gap_analysis_{timestamp}.xlsx"
                    report_path = os.path.join(self.output_folder, report_filename)
                    
                    # Create gap report
                    self.create_gap_report(coverage_analysis, report_path)
                    
                    # Save JSON analysis
                    json_filename = f"{client_name}_gap_analysis_{timestamp}.json"
                    json_path = os.path.join(self.output_folder, json_filename)
                    with open(json_path, 'w') as f:
                        json.dump(coverage_analysis, f, indent=2)
                        
                    successful_analyses += 1
                    self.logger.info(f"Completed gap analysis for {client_name}")
                    
            except Exception as e:
                self.logger.error(f"Failed to analyze {excel_file}: {str(e)}")
                
        self.logger.info(f"Successfully analyzed {successful_analyses}/{len(excel_files)} clients")
        
        # Create analysis summary
        self._create_analysis_summary(successful_analyses, len(excel_files), input_folder)
        
        return self.output_folder
        
    def _create_analysis_summary(self, successful: int, total: int, input_folder: str):
        """Create summary of gap analysis operation."""
        summary = {
            "analysis_summary": {
                "timestamp": datetime.now().isoformat(),
                "input_folder": input_folder,
                "output_folder": self.output_folder,
                "total_excel_files": total,
                "successful_analyses": successful,
                "failed_analyses": total - successful
            },
            "analysis_files": {}
        }
        
        # Add details for each analysis file
        for json_file in Path(self.output_folder).glob("*_gap_analysis_*.json"):
            try:
                with open(json_file, 'r') as f:
                    analysis_data = json.load(f)
                    
                summary["analysis_files"][json_file.name] = {
                    "client_name": analysis_data.get("client_name", "Unknown"),
                    "coverage_percentage": analysis_data.get("coverage_percentage", 0),
                    "covered_techniques": analysis_data.get("covered_techniques", 0),
                    "uncovered_techniques": analysis_data.get("uncovered_techniques", 0),
                    "file_size_mb": round(json_file.stat().st_size / (1024*1024), 3)
                }
            except:
                continue
                
        # Save summary
        summary_path = os.path.join(self.output_folder, "analysis_summary.json")
        with open(summary_path, 'w') as f:
            json.dump(summary, f, indent=2)
            
        self.logger.info(f"Created analysis summary: {summary_path}")