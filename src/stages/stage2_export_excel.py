"""
Stage 2: Export Rules to Excel
Converts JSON rule files to well-formatted Excel files, focusing on enabled rules.
"""

import json
import os
import pandas as pd
from datetime import datetime
from pathlib import Path
from typing import List, Dict


class ExcelExporter:
    """Exports detection rules from JSON to Excel format."""
    
    def __init__(self, base_output_dir, logger):
        """
        Initialize Excel exporter.
        
        Args:
            base_output_dir (str): Base output directory
            logger: Logger instance
        """
        self.base_output_dir = base_output_dir
        self.logger = logger
        self.output_folder = os.path.join(base_output_dir, "02_excel_exports")
        
        # Define the columns for the Excel export
        self.excel_columns = [
            "Rule Name",
            "Description", 
            "Severity",
            "MITRE Tactic",
            "MITRE Technique",
            "Technique ID",
            "MITRE Sub-technique",
            "Sub-technique ID",
            "Tags",
            "MITRE Framework",
            "Creation Date",
            "Rule Type",
            "Platforms",
            "Confidence",
            "False Positive Rate"
        ]
        
    def load_client_rules(self, json_file_path: str) -> Dict:
        """
        Load client rules from JSON file.
        
        Args:
            json_file_path (str): Path to JSON file
            
        Returns:
            dict: Loaded rules data
        """
        try:
            with open(json_file_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            self.logger.error(f"Failed to load {json_file_path}: {str(e)}")
            return {}
            
    def filter_enabled_rules(self, rules_data: Dict) -> List[Dict]:
        """
        Filter only enabled rules from the dataset.
        
        Args:
            rules_data (dict): Complete rules data
            
        Returns:
            list: List of enabled rules
        """
        all_rules = rules_data.get("rules", [])
        enabled_rules = [rule for rule in all_rules if rule.get("enabled", False)]
        
        self.logger.info(f"Filtered {len(enabled_rules)} enabled rules from {len(all_rules)} total rules")
        return enabled_rules
        
    def transform_rules_for_excel(self, rules: List[Dict], client_metadata: Dict) -> pd.DataFrame:
        """
        Transform rules data into Excel-ready format.
        
        Args:
            rules (list): List of rule dictionaries
            client_metadata (dict): Client metadata
            
        Returns:
            pd.DataFrame: Transformed data ready for Excel export
        """
        excel_data = []
        
        for rule in rules:
            # Transform each rule to match Excel columns
            excel_row = {
                "Rule Name": rule.get("name", ""),
                "Description": rule.get("description", ""),
                "Severity": rule.get("severity", ""),
                "MITRE Tactic": rule.get("mitre_tactic", ""),
                "MITRE Technique": rule.get("mitre_technique", ""),
                "Technique ID": rule.get("mitre_technique_id", ""),
                "MITRE Sub-technique": rule.get("mitre_subtechnique", ""),
                "Sub-technique ID": rule.get("mitre_subtechnique_id", ""),
                "Tags": ", ".join(rule.get("tags", [])) if isinstance(rule.get("tags"), list) else str(rule.get("tags", "")),
                "MITRE Framework": rule.get("mitre_framework", "Enterprise"),
                "Creation Date": rule.get("created_date", ""),
                "Rule Type": rule.get("rule_type", ""),
                "Platforms": ", ".join(rule.get("platforms", [])) if isinstance(rule.get("platforms"), list) else str(rule.get("platforms", "")),
                "Confidence": rule.get("confidence", ""),
                "False Positive Rate": rule.get("false_positive_rate", "")
            }
            excel_data.append(excel_row)
            
        # Create DataFrame
        df = pd.DataFrame(excel_data, columns=self.excel_columns)
        
        # Sort by MITRE Tactic, then by Technique ID
        df = df.sort_values(["MITRE Tactic", "Technique ID"], na_position='last')
        
        return df
        
    def create_excel_file(self, df: pd.DataFrame, client_metadata: Dict, output_path: str):
        """
        Create formatted Excel file with multiple sheets.
        
        Args:
            df (pd.DataFrame): Rules data
            client_metadata (dict): Client metadata
            output_path (str): Output file path
        """
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # Main sheet with rules
            df.to_excel(writer, sheet_name='Detection Rules', index=False)
            
            # Summary sheet
            summary_data = self._create_summary_sheet(df, client_metadata)
            summary_df = pd.DataFrame(summary_data)
            summary_df.to_excel(writer, sheet_name='Summary', index=False, header=False)
            
            # MITRE Coverage sheet
            coverage_data = self._create_mitre_coverage_sheet(df)
            coverage_df = pd.DataFrame(coverage_data)
            coverage_df.to_excel(writer, sheet_name='MITRE Coverage', index=False)
            
            # Format the workbook
            self._format_excel_workbook(writer, df)
            
    def _create_summary_sheet(self, df: pd.DataFrame, metadata: Dict) -> List[List]:
        """Create summary information sheet."""
        summary = [
            ["S-Rank Core - Detection Rules Export"],
            [""],
            ["Client Information", ""],
            ["Client Name", metadata.get("client_name", "Unknown")],
            ["Client ID", metadata.get("client_id", "Unknown")],
            ["Export Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            [""],
            ["Rule Statistics", ""],
            ["Total Enabled Rules", len(df)],
            ["Unique MITRE Techniques", df['Technique ID'].nunique()],
            ["Unique MITRE Tactics", df['MITRE Tactic'].nunique()],
            [""],
            ["Severity Distribution", ""],
        ]
        
        # Add severity distribution
        severity_counts = df['Severity'].value_counts()
        for severity, count in severity_counts.items():
            summary.append([f"  {severity}", count])
            
        summary.extend([
            [""],
            ["MITRE Tactic Distribution", ""],
        ])
        
        # Add tactic distribution
        tactic_counts = df['MITRE Tactic'].value_counts()
        for tactic, count in tactic_counts.items():
            summary.append([f"  {tactic}", count])
            
        return summary
        
    def _create_mitre_coverage_sheet(self, df: pd.DataFrame) -> List[Dict]:
        """Create MITRE coverage analysis sheet."""
        coverage_data = []
        
        # Group by MITRE Tactic and count techniques
        tactic_coverage = df.groupby('MITRE Tactic').agg({
            'Technique ID': 'nunique',
            'Rule Name': 'count'
        }).reset_index()
        
        tactic_coverage.columns = ['MITRE Tactic', 'Unique Techniques', 'Total Rules']
        
        # Calculate coverage percentages (this would be more accurate with actual MITRE data)
        known_tactic_totals = {
            'Initial Access': 9,
            'Execution': 12,
            'Persistence': 19,
            'Privilege Escalation': 13,
            'Defense Evasion': 40,
            'Credential Access': 15,
            'Discovery': 29,
            'Lateral Movement': 9,
            'Collection': 17,
            'Command And Control': 16,
            'Exfiltration': 9,
            'Impact': 13
        }
        
        coverage_with_percentages = []
        for _, row in tactic_coverage.iterrows():
            tactic = row['MITRE Tactic']
            covered = row['Unique Techniques']
            total_known = known_tactic_totals.get(tactic, covered)  # Fallback to covered if unknown
            percentage = (covered / total_known * 100) if total_known > 0 else 0
            
            coverage_with_percentages.append({
                'MITRE Tactic': tactic,
                'Covered Techniques': covered,
                'Total Known Techniques': total_known,
                'Coverage Percentage': f"{percentage:.1f}%",
                'Total Rules': row['Total Rules']
            })
            
        return coverage_with_percentages
        
    def _format_excel_workbook(self, writer, df: pd.DataFrame):
        """Apply formatting to the Excel workbook."""
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        
        workbook = writer.book
        
        # Format main rules sheet
        worksheet = writer.sheets['Detection Rules']
        
        # Header formatting
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        for col_num, column in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
                    
            adjusted_width = min(max_length + 2, 50)  # Cap at 50 chars
            worksheet.column_dimensions[column_letter].width = adjusted_width
            
        # Format summary sheet
        summary_ws = writer.sheets['Summary']
        title_font = Font(bold=True, size=14)
        summary_ws['A1'].font = title_font
        
        # Format coverage sheet
        coverage_ws = writer.sheets['MITRE Coverage']
        for col_num in range(1, 6):  # 5 columns in coverage sheet
            cell = coverage_ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    def export_client_rules(self, json_file_path: str) -> str:
        """
        Export a single client's rules to Excel.
        
        Args:
            json_file_path (str): Path to client JSON file
            
        Returns:
            str: Path to created Excel file
        """
        # Load rules data
        rules_data = self.load_client_rules(json_file_path)
        if not rules_data:
            raise ValueError(f"Could not load rules from {json_file_path}")
            
        # Filter enabled rules
        enabled_rules = self.filter_enabled_rules(rules_data)
        if not enabled_rules:
            self.logger.warning(f"No enabled rules found in {json_file_path}")
            
        # Transform to Excel format
        df = self.transform_rules_for_excel(enabled_rules, rules_data.get("metadata", {}))
        
        # Generate output filename
        client_name = rules_data.get("metadata", {}).get("client_id", "unknown")
        timestamp = datetime.now().strftime("%Y%m%d")
        excel_filename = f"{client_name}_enabled_rules_{timestamp}.xlsx"
        excel_path = os.path.join(self.output_folder, excel_filename)
        
        # Ensure output directory exists
        os.makedirs(self.output_folder, exist_ok=True)
        
        # Create Excel file
        self.create_excel_file(df, rules_data.get("metadata", {}), excel_path)
        
        self.logger.info(f"Exported {len(enabled_rules)} enabled rules to {excel_path}")
        return excel_path
        
    def export_all_clients(self, input_folder: str) -> str:
        """
        Export all client JSON files to Excel format.
        
        Args:
            input_folder (str): Folder containing client JSON files
            
        Returns:
            str: Path to output folder containing Excel files
        """
        self.logger.info(f"Starting Excel export from {input_folder}")
        
        # Find all JSON files (excluding summary files)
        json_files = []
        for file_path in Path(input_folder).glob("*.json"):
            if not file_path.name.startswith("fetch_summary"):
                json_files.append(str(file_path))
                
        if not json_files:
            self.logger.warning(f"No client JSON files found in {input_folder}")
            return self.output_folder
            
        successful_exports = 0
        
        for json_file in json_files:
            try:
                excel_path = self.export_client_rules(json_file)
                successful_exports += 1
                self.logger.info(f"Successfully exported {os.path.basename(json_file)}")
                
            except Exception as e:
                self.logger.error(f"Failed to export {json_file}: {str(e)}")
                
        self.logger.info(f"Successfully exported {successful_exports}/{len(json_files)} client files")
        
        # Create export summary
        self._create_export_summary(successful_exports, len(json_files), input_folder)
        
        return self.output_folder
        
    def _create_export_summary(self, successful: int, total: int, input_folder: str):
        """Create summary of the export operation."""
        summary = {
            "export_summary": {
                "timestamp": datetime.now().isoformat(),
                "input_folder": input_folder,
                "output_folder": self.output_folder,
                "total_json_files": total,
                "successful_exports": successful,
                "failed_exports": total - successful
            },
            "exported_files": {}
        }
        
        # Add details for each Excel file
        for excel_file in Path(self.output_folder).glob("*.xlsx"):
            summary["exported_files"][excel_file.name] = {
                "file_size_mb": round(excel_file.stat().st_size / (1024*1024), 2),
                "created": datetime.fromtimestamp(excel_file.stat().st_mtime).isoformat()
            }
            
        # Save summary as JSON
        summary_path = os.path.join(self.output_folder, "export_summary.json")
        with open(summary_path, 'w') as f:
            json.dump(summary, f, indent=2)
            
        self.logger.info(f"Created export summary: {summary_path}")
        
    def get_export_statistics(self, excel_file_path: str) -> Dict:
        """
        Get statistics about an exported Excel file.
        
        Args:
            excel_file_path (str): Path to Excel file
            
        Returns:
            dict: Statistics about the export
        """
        try:
            # Read the main sheet
            df = pd.read_excel(excel_file_path, sheet_name='Detection Rules')
            
            stats = {
                "file_path": excel_file_path,
                "total_rules": len(df),
                "unique_techniques": df['Technique ID'].nunique(),
                "unique_tactics": df['MITRE Tactic'].nunique(),
                "severity_distribution": df['Severity'].value_counts().to_dict(),
                "tactic_distribution": df['MITRE Tactic'].value_counts().to_dict(),
                "file_size_mb": round(os.path.getsize(excel_file_path) / (1024*1024), 2)
            }
            
            return stats
            
        except Exception as e:
            self.logger.error(f"Failed to get statistics for {excel_file_path}: {str(e)}")
            return {}