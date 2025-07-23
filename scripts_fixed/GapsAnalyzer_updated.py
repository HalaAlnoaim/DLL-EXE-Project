#!/usr/bin/env python3
"""
S-Rank Core - Stage 3: Enhanced Gap Analysis (Updated Version)
Now with ConfigManager integration and enhanced technique processing.
"""

import os
import json
import sys
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Optional, Set
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
import logging
import requests

# Add the project root to Python path
project_root = Path(__file__).parent.parent
sys.path.append(str(project_root))

from config_manager import ConfigManager

class Color:
    """ANSI color codes for console output."""
    OKGREEN = "\033[92m"
    WARNING = "\033[93m"
    FAIL = "\033[91m"
    ENDC = "\033[0m"
    BOLD = "\033[1m"
    CYAN = "\033[96m"

class EnhancedGapsAnalyzer:
    """Enhanced Gap Analyzer with ConfigManager integration and technique processing."""
    
    def __init__(self):
        self.setup_logging()
        self.config_manager = ConfigManager()
        self.validate_environment()
        self.load_configuration()
        self.mitre_data = None
        
    def setup_logging(self):
        """Setup logging configuration."""
        log_level = os.environ.get("SRANK_LOG_LEVEL", "INFO")
        logging.basicConfig(
            level=getattr(logging, log_level.upper()),
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[logging.StreamHandler(sys.stdout)]
        )
        self.logger = logging.getLogger(__name__)
        
    def validate_environment(self):
        """Validate environment variables and configuration."""
        # Get report directory from environment
        self.report_dir = os.environ.get("SRANK_REPORT_DIR")
        if not self.report_dir:
            raise ValueError("SRANK_REPORT_DIR not found. Please run this script via main.py")
            
        # Get pipeline timestamp
        self.pipeline_timestamp = os.environ.get("SRANK_PIPELINE_TIMESTAMP")
        if not self.pipeline_timestamp:
            self.logger.warning("SRANK_PIPELINE_TIMESTAMP not found, using current timestamp")
            self.pipeline_timestamp = datetime.now().strftime("%Y_%m_%d_%H%M")
            
        self.report_dir = Path(self.report_dir)
        
        # Find the latest Exported Rules folder
        exported_folders = sorted(self.report_dir.glob("Exported_Rules_*"))
        if not exported_folders:
            raise FileNotFoundError("No Exported_Rules_* folder found. Please run Stage 2 (Export to Excel) first.")
        self.latest_exported_folder = exported_folders[-1]
        
        # Create gaps analysis folder
        self.gaps_dir = self.report_dir / f"Reports_{self.pipeline_timestamp}"
        self.gaps_dir.mkdir(parents=True, exist_ok=True)
        
        self.logger.info(f"Input folder: {self.latest_exported_folder}")
        self.logger.info(f"Output folder: {self.gaps_dir}")
        
    def load_configuration(self):
        """Load configuration from ConfigManager."""
        # Get processing settings
        self.processing_settings = self.config_manager.get_processing_settings()
        
        # Enhanced column styling
        self.header_fill = PatternFill(start_color="702386", end_color="702386", fill_type="solid")
        self.covered_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")  # Light green
        self.gap_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")     # Light red
        self.header_font = Font(bold=True, color="FFFFFF")
        self.default_font = Font(color="000000")
        
        self.logger.info("Loaded enhanced configuration for gap analysis")
        
    def load_mitre_attack_data(self) -> bool:
        """Load MITRE ATT&CK framework data."""
        try:
            print(f"{Color.CYAN}📥 Loading MITRE ATT&CK framework data...{Color.ENDC}")
            
            # Try to load from local cache first
            cache_file = self.report_dir / "mitre_attack_cache.json"
            if cache_file.exists():
                try:
                    with open(cache_file, 'r', encoding='utf-8') as f:
                        self.mitre_data = json.load(f)
                    self.logger.info("Loaded MITRE data from cache")
                    print(f"{Color.OKGREEN}✅ MITRE data loaded from cache{Color.ENDC}")
                    return True
                except Exception as e:
                    self.logger.warning(f"Failed to load cache: {e}")
            
            # Download fresh data from MITRE
            print(f"{Color.CYAN}🌐 Downloading fresh MITRE ATT&CK data...{Color.ENDC}")
            
            mitre_url = "https://raw.githubusercontent.com/mitre/cti/master/enterprise-attack/enterprise-attack.json"
            
            try:
                response = requests.get(mitre_url, timeout=30)
                response.raise_for_status()
                
                mitre_raw = response.json()
                
                # Process MITRE data
                self.mitre_data = {
                    "techniques": {},
                    "tactics": {},
                    "subtechniques": {},
                    "metadata": {
                        "loaded_at": datetime.now().isoformat(),
                        "version": mitre_raw.get("spec_version", "unknown"),
                        "source": "MITRE CTI Repository"
                    }
                }
                
                # Extract techniques, subtechniques, and tactics
                for obj in mitre_raw.get("objects", []):
                    if obj.get("type") == "attack-pattern":
                        # Extract technique ID and name
                        external_refs = obj.get("external_references", [])
                        mitre_id = None
                        for ref in external_refs:
                            if ref.get("source_name") == "mitre-attack":
                                mitre_id = ref.get("external_id")
                                break
                        
                        if mitre_id:
                            tech_data = {
                                "id": mitre_id,
                                "name": obj.get("name", ""),
                                "description": obj.get("description", ""),
                                "tactics": [phase.replace("-", "_") for phase in obj.get("kill_chain_phases", [])],
                                "platforms": obj.get("x_mitre_platforms", []),
                                "is_subtechnique": "." in mitre_id
                            }
                            
                            if "." in mitre_id:
                                # This is a subtechnique
                                self.mitre_data["subtechniques"][mitre_id] = tech_data
                            else:
                                # This is a main technique
                                self.mitre_data["techniques"][mitre_id] = tech_data
                    
                    elif obj.get("type") == "x-mitre-tactic":
                        # Extract tactic data
                        external_refs = obj.get("external_references", [])
                        tactic_id = None
                        for ref in external_refs:
                            if ref.get("source_name") == "mitre-attack":
                                tactic_id = ref.get("external_id")
                                break
                        
                        if tactic_id:
                            self.mitre_data["tactics"][tactic_id] = {
                                "id": tactic_id,
                                "name": obj.get("name", ""),
                                "description": obj.get("description", ""),
                                "shortname": obj.get("x_mitre_shortname", "")
                            }
                
                # Save to cache
                with open(cache_file, 'w', encoding='utf-8') as f:
                    json.dump(self.mitre_data, f, indent=2, ensure_ascii=False)
                
                technique_count = len(self.mitre_data["techniques"])
                subtechnique_count = len(self.mitre_data["subtechniques"])
                tactic_count = len(self.mitre_data["tactics"])
                
                print(f"{Color.OKGREEN}✅ MITRE data loaded: {technique_count} techniques, {subtechnique_count} subtechniques, {tactic_count} tactics{Color.ENDC}")
                self.logger.info(f"MITRE data loaded successfully")
                return True
                
            except requests.RequestException as e:
                self.logger.error(f"Failed to download MITRE data: {e}")
                print(f"{Color.FAIL}❌ Failed to download MITRE data. Using fallback method.{Color.ENDC}")
                return self.create_fallback_mitre_data()
                
        except Exception as e:
            self.logger.error(f"Error loading MITRE data: {e}")
            print(f"{Color.FAIL}❌ Error loading MITRE data: {e}{Color.ENDC}")
            return self.create_fallback_mitre_data()
            
    def create_fallback_mitre_data(self) -> bool:
        """Create minimal fallback MITRE data structure."""
        try:
            self.logger.warning("Creating fallback MITRE data")
            
            # Basic MITRE structure for common techniques
            fallback_techniques = {
                "T1055": {"id": "T1055", "name": "Process Injection", "tactics": ["defense-evasion", "privilege-escalation"]},
                "T1059": {"id": "T1059", "name": "Command and Scripting Interpreter", "tactics": ["execution"]},
                "T1071": {"id": "T1071", "name": "Application Layer Protocol", "tactics": ["command-and-control"]},
                "T1083": {"id": "T1083", "name": "File and Directory Discovery", "tactics": ["discovery"]},
                "T1087": {"id": "T1087", "name": "Account Discovery", "tactics": ["discovery"]},
                "T1090": {"id": "T1090", "name": "Proxy", "tactics": ["command-and-control"]},
                "T1105": {"id": "T1105", "name": "Ingress Tool Transfer", "tactics": ["command-and-control"]},
                "T1134": {"id": "T1134", "name": "Access Token Manipulation", "tactics": ["defense-evasion", "privilege-escalation"]},
                "T1190": {"id": "T1190", "name": "Exploit Public-Facing Application", "tactics": ["initial-access"]},
                "T1566": {"id": "T1566", "name": "Phishing", "tactics": ["initial-access"]}
            }
            
            fallback_tactics = {
                "TA0001": {"id": "TA0001", "name": "Initial Access", "shortname": "initial-access"},
                "TA0002": {"id": "TA0002", "name": "Execution", "shortname": "execution"},
                "TA0003": {"id": "TA0003", "name": "Persistence", "shortname": "persistence"},
                "TA0004": {"id": "TA0004", "name": "Privilege Escalation", "shortname": "privilege-escalation"},
                "TA0005": {"id": "TA0005", "name": "Defense Evasion", "shortname": "defense-evasion"},
                "TA0006": {"id": "TA0006", "name": "Credential Access", "shortname": "credential-access"},
                "TA0007": {"id": "TA0007", "name": "Discovery", "shortname": "discovery"},
                "TA0008": {"id": "TA0008", "name": "Lateral Movement", "shortname": "lateral-movement"},
                "TA0009": {"id": "TA0009", "name": "Collection", "shortname": "collection"},
                "TA0011": {"id": "TA0011", "name": "Command and Control", "shortname": "command-and-control"},
                "TA0010": {"id": "TA0010", "name": "Exfiltration", "shortname": "exfiltration"},
                "TA0040": {"id": "TA0040", "name": "Impact", "shortname": "impact"}
            }
            
            self.mitre_data = {
                "techniques": fallback_techniques,
                "tactics": fallback_tactics,
                "subtechniques": {},
                "metadata": {
                    "loaded_at": datetime.now().isoformat(),
                    "version": "fallback",
                    "source": "Fallback Basic Data"
                }
            }
            
            print(f"{Color.WARNING}⚠️  Using fallback MITRE data with {len(fallback_techniques)} basic techniques{Color.ENDC}")
            return True
            
        except Exception as e:
            self.logger.error(f"Failed to create fallback data: {e}")
            return False
            
    def analyze_client_coverage(self, excel_file: Path) -> Optional[Dict]:
        """Analyze MITRE coverage for a single client Excel file."""
        try:
            # Extract client ID from filename
            client_id = excel_file.stem.split('_')[0]
            
            # Get client config
            client_config = self.config_manager.get_client_config(client_id)
            client_name = client_config.get('name', client_id) if client_config else client_id
            
            self.logger.info(f"Analyzing coverage for {client_name} ({client_id})")
            
            # Read Excel file
            try:
                df = pd.read_excel(excel_file, sheet_name="Enabled Rules", header=1)  # Skip summary row
            except Exception as e:
                self.logger.error(f"Failed to read Excel file {excel_file}: {e}")
                return None
            
            # Extract techniques from the data
            covered_techniques = set()
            covered_subtechniques = set()
            
            # Process technique IDs
            for _, row in df.iterrows():
                # Main techniques
                tech_ids = str(row.get("TECHNIQUE ID", "")).strip()
                if tech_ids and tech_ids != "nan":
                    for tech_id in tech_ids.split(", "):
                        tech_id = tech_id.strip()
                        if tech_id and "." not in tech_id:  # Main technique
                            covered_techniques.add(tech_id)
                
                # Subtechniques
                subtech_ids = str(row.get("SUBTECHNIQUE ID", "")).strip()
                if subtech_ids and subtech_ids != "nan":
                    for subtech_id in subtech_ids.split(", "):
                        subtech_id = subtech_id.strip()
                        if subtech_id:
                            covered_subtechniques.add(subtech_id)
                            # Also add the parent technique
                            parent_tech = subtech_id.split('.')[0]
                            covered_techniques.add(parent_tech)
                
                # Sub-subtechniques
                sub_subtech_ids = str(row.get("SUB-SUBTECHNIQUE ID", "")).strip()
                if sub_subtech_ids and sub_subtech_ids != "nan":
                    for sub_subtech_id in sub_subtech_ids.split(", "):
                        sub_subtech_id = sub_subtech_id.strip()
                        if sub_subtech_id:
                            covered_subtechniques.add(sub_subtech_id)
                            # Also add the parent technique
                            parent_tech = sub_subtech_id.split('.')[0]
                            covered_techniques.add(parent_tech)
            
            # Calculate gaps
            all_mitre_techniques = set(self.mitre_data["techniques"].keys())
            all_mitre_subtechniques = set(self.mitre_data["subtechniques"].keys())
            
            technique_gaps = all_mitre_techniques - covered_techniques
            subtechnique_gaps = all_mitre_subtechniques - covered_subtechniques
            
            # Calculate coverage by tactic
            tactic_coverage = {}
            for tactic_id, tactic_data in self.mitre_data["tactics"].items():
                tactic_name = tactic_data["name"]
                tactic_shortname = tactic_data.get("shortname", tactic_name.lower().replace(" ", "-"))
                
                # Find techniques for this tactic
                tactic_techniques = set()
                for tech_id, tech_data in self.mitre_data["techniques"].items():
                    if tactic_shortname in tech_data.get("tactics", []):
                        tactic_techniques.add(tech_id)
                
                # Calculate coverage
                covered_in_tactic = tactic_techniques & covered_techniques
                coverage_percentage = (len(covered_in_tactic) / len(tactic_techniques) * 100) if tactic_techniques else 0
                
                tactic_coverage[tactic_name] = {
                    "total_techniques": len(tactic_techniques),
                    "covered_techniques": len(covered_in_tactic),
                    "coverage_percentage": coverage_percentage,
                    "gaps": tactic_techniques - covered_techniques
                }
            
            return {
                "client_id": client_id,
                "client_name": client_name,
                "total_rules": len(df),
                "covered_techniques": covered_techniques,
                "covered_subtechniques": covered_subtechniques,
                "technique_gaps": technique_gaps,
                "subtechnique_gaps": subtechnique_gaps,
                "coverage_stats": {
                    "techniques_covered": len(covered_techniques),
                    "techniques_total": len(all_mitre_techniques),
                    "techniques_coverage_pct": (len(covered_techniques) / len(all_mitre_techniques) * 100) if all_mitre_techniques else 0,
                    "subtechniques_covered": len(covered_subtechniques),
                    "subtechniques_total": len(all_mitre_subtechniques),
                    "subtechniques_coverage_pct": (len(covered_subtechniques) / len(all_mitre_subtechniques) * 100) if all_mitre_subtechniques else 0
                },
                "tactic_coverage": tactic_coverage,
                "analysis_timestamp": datetime.now().isoformat()
            }
            
        except Exception as e:
            self.logger.error(f"Failed to analyze {excel_file}: {e}")
            return None
            
    def create_gaps_report(self, analysis_data: Dict, output_path: Path) -> bool:
        """Create enhanced gaps analysis Excel report."""
        try:
            wb = Workbook()
            
            # Sheet 1: Coverage Overview
            self.create_coverage_overview_sheet(wb, analysis_data)
            
            # Sheet 2: Technique Gaps
            self.create_technique_gaps_sheet(wb, analysis_data)
            
            # Sheet 3: Tactic Coverage
            self.create_tactic_coverage_sheet(wb, analysis_data)
            
            # Sheet 4: Recommendations
            self.create_recommendations_sheet(wb, analysis_data)
            
            # Save workbook
            wb.save(output_path)
            
            # Validate file creation
            if output_path.exists() and output_path.stat().st_size > 0:
                self.logger.info(f"Gaps analysis report created: {output_path}")
                return True
            else:
                self.logger.error(f"Failed to create gaps report: {output_path}")
                return False
                
        except Exception as e:
            self.logger.error(f"Failed to create gaps report: {e}")
            return False
            
    def create_coverage_overview_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create coverage overview sheet."""
        ws = wb.active
        ws.title = "Coverage Overview"
        
        client_name = analysis_data["client_name"]
        stats = analysis_data["coverage_stats"]
        
        # Title and client info
        overview_data = [
            [f"MITRE ATT&CK Coverage Analysis - {client_name}", ""],
            ["", ""],
            ["Analysis Summary", ""],
            ["Client Name", client_name],
            ["Analysis Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Total Detection Rules", analysis_data["total_rules"]],
            ["", ""],
            ["Technique Coverage", ""],
            ["Techniques Covered", f"{stats['techniques_covered']}/{stats['techniques_total']} ({stats['techniques_coverage_pct']:.1f}%)"],
            ["Subtechniques Covered", f"{stats['subtechniques_covered']}/{stats['subtechniques_total']} ({stats['subtechniques_coverage_pct']:.1f}%)"],
            ["", ""],
            ["Top Coverage Tactics", ""],
        ]
        
        # Add top tactics by coverage
        tactic_coverage = analysis_data["tactic_coverage"]
        sorted_tactics = sorted(tactic_coverage.items(), key=lambda x: x[1]["coverage_percentage"], reverse=True)
        
        for tactic_name, tactic_data in sorted_tactics[:10]:  # Top 10
            coverage_pct = tactic_data["coverage_percentage"]
            covered = tactic_data["covered_techniques"]
            total = tactic_data["total_techniques"]
            overview_data.append([f"  {tactic_name}", f"{covered}/{total} ({coverage_pct:.1f}%)"])
        
        # Write data to sheet
        for row_idx, row_data in enumerate(overview_data, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Style headers
                if col_idx == 1 and value and not str(value).startswith("  "):
                    if row_idx == 1:  # Main title
                        cell.font = Font(bold=True, size=16)
                    else:
                        cell.font = Font(bold=True)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 30
        
    def create_technique_gaps_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create technique gaps analysis sheet."""
        ws = wb.create_sheet(title="Technique Gaps")
        
        # Headers
        headers = ["Technique ID", "Technique Name", "Tactics", "Status", "Priority"]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add all techniques with coverage status
        row_idx = 2
        
        # Sort techniques by ID
        all_techniques = list(self.mitre_data["techniques"].keys())
        all_techniques.sort()
        
        for tech_id in all_techniques:
            tech_data = self.mitre_data["techniques"][tech_id]
            is_covered = tech_id in analysis_data["covered_techniques"]
            
            # Determine priority based on tactics
            tactics_list = tech_data.get("tactics", [])
            priority = "High" if any(tactic in ["initial-access", "execution", "persistence"] for tactic in tactics_list) else "Medium"
            
            # Write row data
            ws.cell(row=row_idx, column=1, value=tech_id)
            ws.cell(row=row_idx, column=2, value=tech_data["name"])
            ws.cell(row=row_idx, column=3, value=", ".join(tactics_list))
            
            status_cell = ws.cell(row=row_idx, column=4, value="Covered" if is_covered else "Gap")
            status_cell.fill = self.covered_fill if is_covered else self.gap_fill
            
            priority_cell = ws.cell(row=row_idx, column=5, value=priority if not is_covered else "N/A")
            
            row_idx += 1
        
        # Auto-adjust column widths
        widths = [15, 40, 30, 15, 15]
        for col_idx, width in enumerate(widths, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
            
        # Add auto-filter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_idx-1}"
        
    def create_tactic_coverage_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create tactic coverage analysis sheet."""
        ws = wb.create_sheet(title="Tactic Coverage")
        
        # Headers
        headers = ["Tactic", "Total Techniques", "Covered", "Coverage %", "Gap Count"]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add tactic data
        tactic_coverage = analysis_data["tactic_coverage"]
        sorted_tactics = sorted(tactic_coverage.items(), key=lambda x: x[1]["coverage_percentage"], reverse=True)
        
        for row_idx, (tactic_name, tactic_data) in enumerate(sorted_tactics, start=2):
            ws.cell(row=row_idx, column=1, value=tactic_name)
            ws.cell(row=row_idx, column=2, value=tactic_data["total_techniques"])
            ws.cell(row=row_idx, column=3, value=tactic_data["covered_techniques"])
            
            coverage_pct = tactic_data["coverage_percentage"]
            coverage_cell = ws.cell(row=row_idx, column=4, value=f"{coverage_pct:.1f}%")
            
            # Color code based on coverage
            if coverage_pct >= 80:
                coverage_cell.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
            elif coverage_pct >= 50:
                coverage_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            else:
                coverage_cell.fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
            
            gap_count = len(tactic_data["gaps"])
            ws.cell(row=row_idx, column=5, value=gap_count)
        
        # Auto-adjust column widths
        widths = [25, 18, 15, 15, 15]
        for col_idx, width in enumerate(widths, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
            
    def create_recommendations_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create recommendations sheet."""
        ws = wb.create_sheet(title="Recommendations")
        
        client_name = analysis_data["client_name"]
        stats = analysis_data["coverage_stats"]
        
        # Generate recommendations based on gaps
        recommendations = [
            [f"Detection Coverage Recommendations - {client_name}", ""],
            ["", ""],
            ["Executive Summary", ""],
            [f"Current technique coverage: {stats['techniques_coverage_pct']:.1f}%", ""],
            [f"Current subtechnique coverage: {stats['subtechniques_coverage_pct']:.1f}%", ""],
            ["", ""],
            ["Priority Recommendations", ""],
        ]
        
        # Analyze gaps and provide recommendations
        tactic_coverage = analysis_data["tactic_coverage"]
        
        # Find tactics with lowest coverage
        low_coverage_tactics = [
            (name, data) for name, data in tactic_coverage.items() 
            if data["coverage_percentage"] < 50 and data["total_techniques"] > 0
        ]
        low_coverage_tactics.sort(key=lambda x: x[1]["coverage_percentage"])
        
        if low_coverage_tactics:
            recommendations.append(["1. Focus on Low Coverage Tactics:", ""])
            for i, (tactic_name, tactic_data) in enumerate(low_coverage_tactics[:5], 1):
                gap_count = len(tactic_data["gaps"])
                coverage_pct = tactic_data["coverage_percentage"]
                recommendations.append([
                    f"   {i}. {tactic_name}: {coverage_pct:.1f}% coverage, {gap_count} gaps",
                    "High Priority"
                ])
        
        # Add general recommendations
        recommendations.extend([
            ["", ""],
            ["2. General Recommendations:", ""],
            ["   • Implement detection for Initial Access techniques", "High Priority"],
            ["   • Enhance Command and Control monitoring", "Medium Priority"],
            ["   • Improve Persistence detection capabilities", "High Priority"],
            ["   • Deploy behavioral analytics for Defense Evasion", "Medium Priority"],
            ["   • Strengthen Privilege Escalation monitoring", "High Priority"],
            ["", ""],
            ["3. Implementation Approach:", ""],
            ["   • Phase 1: Address critical gaps in high-priority tactics", ""],
            ["   • Phase 2: Improve coverage in medium-priority areas", ""],
            ["   • Phase 3: Fine-tune detection accuracy and reduce false positives", ""],
        ])
        
        # Write recommendations to sheet
        for row_idx, row_data in enumerate(recommendations, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Style headers and priorities
                if col_idx == 1 and value and not str(value).startswith("   "):
                    if row_idx == 1:  # Main title
                        cell.font = Font(bold=True, size=16)
                    elif any(char.isdigit() for char in str(value)[:2]):  # Numbered recommendations
                        cell.font = Font(bold=True)
                elif col_idx == 2 and "Priority" in str(value):
                    if "High" in str(value):
                        cell.fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
                    elif "Medium" in str(value):
                        cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 60
        ws.column_dimensions['B'].width = 20
        
    def analyze_all_clients(self) -> bool:
        """Analyze gaps for all client Excel files."""
        print(f"{Color.BOLD}🚀 S-Rank Core - Stage 3: Enhanced Gap Analysis{Color.ENDC}")
        print(f"{Color.CYAN}📋 Configuration Summary:{Color.ENDC}")
        self.config_manager.print_config_summary()
        print(f"\n{Color.CYAN}📂 Input folder:{Color.ENDC} {self.latest_exported_folder}")
        print(f"{Color.CYAN}📂 Output folder:{Color.ENDC} {self.gaps_dir}")
        print()
        
        # Load MITRE ATT&CK data
        if not self.load_mitre_attack_data():
            print(f"{Color.FAIL}❌ Failed to load MITRE ATT&CK data. Cannot proceed with gap analysis.{Color.ENDC}")
            return False
        
        # Find Excel files to analyze
        excel_files = list(self.latest_exported_folder.glob("*_enabled_rules_*.xlsx"))
        
        if not excel_files:
            print(f"{Color.FAIL}❌ No client Excel files found in {self.latest_exported_folder}{Color.ENDC}")
            return False
        
        print(f"{Color.CYAN}📄 Found {len(excel_files)} client files to analyze{Color.ENDC}\n")
        
        successful_analyses = 0
        analysis_summary = []
        
        for excel_file in excel_files:
            try:
                client_id = excel_file.stem.split('_')[0]
                client_config = self.config_manager.get_client_config(client_id)
                client_name = client_config.get('name', client_id) if client_config else client_id
                
                print(f"{Color.CYAN}Analyzing:{Color.ENDC} {client_name} ({excel_file.name})...")
                
                # Perform gap analysis
                analysis_data = self.analyze_client_coverage(excel_file)
                
                if analysis_data:
                    # Generate output filename
                    output_filename = f"{client_id}_gaps_analysis_{self.pipeline_timestamp}.xlsx"
                    output_path = self.gaps_dir / output_filename
                    
                    # Create gaps report
                    if self.create_gaps_report(analysis_data, output_path):
                        coverage_pct = analysis_data["coverage_stats"]["techniques_coverage_pct"]
                        gap_count = len(analysis_data["technique_gaps"])
                        
                        print(f"{Color.OKGREEN}[✓] {client_name}: {coverage_pct:.1f}% coverage, {gap_count} gaps identified{Color.ENDC}")
                        successful_analyses += 1
                        
                        analysis_summary.append({
                            "client_id": client_id,
                            "client_name": client_name,
                            "coverage_percentage": coverage_pct,
                            "gaps_count": gap_count,
                            "total_rules": analysis_data["total_rules"],
                            "output_file": output_filename,
                            "success": True
                        })
                    else:
                        print(f"{Color.FAIL}[✗] {client_name}: Failed to create gaps report{Color.ENDC}")
                        analysis_summary.append({
                            "client_id": client_id,
                            "client_name": client_name,
                            "coverage_percentage": 0,
                            "gaps_count": 0,
                            "total_rules": 0,
                            "output_file": "N/A",
                            "success": False
                        })
                else:
                    print(f"{Color.FAIL}[✗] {client_name}: Failed to analyze coverage{Color.ENDC}")
                    analysis_summary.append({
                        "client_id": client_id,
                        "client_name": client_name,
                        "coverage_percentage": 0,
                        "gaps_count": 0,
                        "total_rules": 0,
                        "output_file": "N/A",
                        "success": False
                    })
                    
            except Exception as e:
                print(f"{Color.FAIL}[✗] {excel_file.name}: Error - {e}{Color.ENDC}")
                self.logger.error(f"Error analyzing {excel_file.name}: {e}")
        
        # Create analysis summary
        self.create_analysis_summary(analysis_summary)
        
        # Final status
        print(f"\n{Color.BOLD}📊 Gap Analysis Summary:{Color.ENDC}")
        print(f"   • Files analyzed: {successful_analyses}/{len(excel_files)}")
        print(f"   • Success rate: {(successful_analyses/len(excel_files)*100):.1f}%")
        
        if analysis_summary:
            avg_coverage = sum(s["coverage_percentage"] for s in analysis_summary if s["success"]) / max(successful_analyses, 1)
            total_gaps = sum(s["gaps_count"] for s in analysis_summary if s["success"])
            print(f"   • Average coverage: {avg_coverage:.1f}%")
            print(f"   • Total gaps identified: {total_gaps}")
        
        if successful_analyses == len(excel_files):
            print(f"{Color.OKGREEN}🎉 All clients analyzed successfully!{Color.ENDC}")
            return True
        elif successful_analyses > 0:
            print(f"{Color.WARNING}⚠️  Partially completed with some errors.{Color.ENDC}")
            return True
        else:
            print(f"{Color.FAIL}❌ Failed to analyze any clients.{Color.ENDC}")
            return False
            
    def create_analysis_summary(self, analysis_summary: List[Dict]):
        """Create enhanced analysis summary."""
        try:
            successful_analyses = [s for s in analysis_summary if s["success"]]
            
            summary_data = {
                "gap_analysis_summary": {
                    "timestamp": datetime.now().isoformat(),
                    "pipeline_timestamp": self.pipeline_timestamp,
                    "processing_version": "2.0_enhanced_gaps",
                    "mitre_data_version": self.mitre_data.get("metadata", {}).get("version", "unknown"),
                    "total_clients": len(analysis_summary),
                    "successful_analyses": len(successful_analyses),
                    "failed_analyses": len(analysis_summary) - len(successful_analyses),
                    "average_coverage": sum(s["coverage_percentage"] for s in successful_analyses) / max(len(successful_analyses), 1),
                    "total_gaps_identified": sum(s["gaps_count"] for s in successful_analyses),
                    "total_rules_analyzed": sum(s["total_rules"] for s in successful_analyses)
                },
                "client_analyses": analysis_summary,
                "configuration": {
                    "technique_naming_format": self.config_manager.get_technique_naming_format(),
                    "processing_settings": self.processing_settings
                },
                "mitre_metadata": self.mitre_data.get("metadata", {})
            }
            
            summary_path = self.gaps_dir / "gap_analysis_summary_enhanced.json"
            with open(summary_path, 'w', encoding='utf-8') as f:
                json.dump(summary_data, f, indent=2, ensure_ascii=False)
                
            self.logger.info(f"Enhanced gap analysis summary saved to: {summary_path}")
            
        except Exception as e:
            self.logger.error(f"Failed to create analysis summary: {e}")


def main():
    """Entry point for the enhanced gap analyzer script."""
    try:
        analyzer = EnhancedGapsAnalyzer()
        success = analyzer.analyze_all_clients()
        return 0 if success else 1
    except KeyboardInterrupt:
        print(f"\n{Color.WARNING}⚠️  Operation cancelled by user.{Color.ENDC}")
        return 130
    except Exception as e:
        print(f"{Color.FAIL}💥 Fatal error: {e}{Color.ENDC}")
        logging.error(f"Fatal error: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())