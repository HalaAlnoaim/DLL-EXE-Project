#!/usr/bin/env python3
"""
S-Rank Core - Stage 4: Enhanced Common Gaps Scanner (Updated Version)
Now with ConfigManager integration and enhanced aggregation analysis.
"""

import os
import json
import sys
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Optional, Set, Tuple
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, PieChart
import logging
from collections import Counter, defaultdict

# Add the project root to Python path
project_root = Path(__file__).parent.parent
sys.path.append(str(project_root))

from config_manager import ConfigManager

class Color:
    """ANSI color codes for console output."""
    OKGREEN = "\033[92m"
    WARNING = "\033[93m"
    FAIL = "\033[91m"
    ENDC = "\033[0m"
    BOLD = "\033[1m"
    CYAN = "\033[96m"

class EnhancedGapsScanner:
    """Enhanced Common Gaps Scanner with ConfigManager integration."""
    
    def __init__(self):
        self.setup_logging()
        self.config_manager = ConfigManager()
        self.validate_environment()
        self.load_configuration()
        
    def setup_logging(self):
        """Setup logging configuration."""
        log_level = os.environ.get("SRANK_LOG_LEVEL", "INFO")
        logging.basicConfig(
            level=getattr(logging, log_level.upper()),
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[logging.StreamHandler(sys.stdout)]
        )
        self.logger = logging.getLogger(__name__)
        
    def validate_environment(self):
        """Validate environment variables and configuration."""
        # Get report directory from environment
        self.report_dir = os.environ.get("SRANK_REPORT_DIR")
        if not self.report_dir:
            raise ValueError("SRANK_REPORT_DIR not found. Please run this script via main.py")
            
        # Get pipeline timestamp
        self.pipeline_timestamp = os.environ.get("SRANK_PIPELINE_TIMESTAMP")
        if not self.pipeline_timestamp:
            self.logger.warning("SRANK_PIPELINE_TIMESTAMP not found, using current timestamp")
            self.pipeline_timestamp = datetime.now().strftime("%Y_%m_%d_%H%M")
            
        self.report_dir = Path(self.report_dir)
        
        # Find the latest Reports folder (gap analysis results)
        reports_folders = sorted(self.report_dir.glob("Reports_*"))
        if not reports_folders:
            raise FileNotFoundError("No Reports_* folder found. Please run Stage 3 (Gap Analysis) first.")
        self.latest_reports_folder = reports_folders[-1]
        
        # Create common gaps scanner folder
        self.common_gaps_dir = self.report_dir / f"CommonGaps_{self.pipeline_timestamp}"
        self.common_gaps_dir.mkdir(parents=True, exist_ok=True)
        
        self.logger.info(f"Input folder: {self.latest_reports_folder}")
        self.logger.info(f"Output folder: {self.common_gaps_dir}")
        
    def load_configuration(self):
        """Load configuration from ConfigManager."""
        # Get processing settings
        self.processing_settings = self.config_manager.get_processing_settings()
        
        # Enhanced styling
        self.header_fill = PatternFill(start_color="702386", end_color="702386", fill_type="solid")
        self.critical_fill = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")     # Red
        self.high_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")        # Orange
        self.medium_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")      # Yellow
        self.low_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")         # Light Green
        self.header_font = Font(bold=True, color="FFFFFF")
        self.default_font = Font(color="000000")
        self.bold_font = Font(bold=True, color="000000")
        
        self.logger.info("Loaded enhanced configuration for common gaps scanner")
        
    def load_gap_analysis_data(self) -> Dict[str, Dict]:
        """Load gap analysis data from all client reports."""
        gap_data = {}
        
        # Look for gap analysis summary file first
        summary_file = self.latest_reports_folder / "gap_analysis_summary_enhanced.json"
        if summary_file.exists():
            try:
                with open(summary_file, 'r', encoding='utf-8') as f:
                    summary_data = json.load(f)
                    
                self.logger.info("Loaded gap analysis summary")
                
                # Extract basic data from summary
                for client_analysis in summary_data.get("client_analyses", []):
                    if client_analysis.get("success"):
                        client_id = client_analysis["client_id"]
                        gap_data[client_id] = {
                            "client_name": client_analysis["client_name"],
                            "coverage_percentage": client_analysis["coverage_percentage"],
                            "gaps_count": client_analysis["gaps_count"],
                            "total_rules": client_analysis["total_rules"],
                            "detailed_data": None  # Will load from individual files
                        }
                        
            except Exception as e:
                self.logger.warning(f"Failed to load summary file: {e}")
        
        # Load detailed data from individual gap analysis files
        gap_files = list(self.latest_reports_folder.glob("*_gaps_analysis_*.xlsx"))
        
        for gap_file in gap_files:
            try:
                client_id = gap_file.stem.split('_')[0]
                
                # Get client config for name
                client_config = self.config_manager.get_client_config(client_id)
                client_name = client_config.get('name', client_id) if client_config else client_id
                
                # Read the Excel file to extract detailed gap information
                try:
                    # Read technique gaps sheet
                    df_gaps = pd.read_excel(gap_file, sheet_name="Technique Gaps")
                    
                    # Read tactic coverage sheet  
                    df_tactics = pd.read_excel(gap_file, sheet_name="Tactic Coverage")
                    
                    # Extract gaps (techniques with "Gap" status)
                    gaps = []
                    for _, row in df_gaps.iterrows():
                        if row.get("Status") == "Gap":
                            gaps.append({
                                "technique_id": row.get("Technique ID", ""),
                                "technique_name": row.get("Technique Name", ""),
                                "tactics": row.get("Tactics", ""),
                                "priority": row.get("Priority", "Medium")
                            })
                    
                    # Extract tactic coverage data
                    tactic_coverage = {}
                    for _, row in df_tactics.iterrows():
                        tactic_name = row.get("Tactic", "")
                        if tactic_name:
                            coverage_str = str(row.get("Coverage %", "0%")).replace("%", "")
                            try:
                                coverage_pct = float(coverage_str)
                            except ValueError:
                                coverage_pct = 0.0
                                
                            tactic_coverage[tactic_name] = {
                                "total_techniques": row.get("Total Techniques", 0),
                                "covered": row.get("Covered", 0),
                                "coverage_percentage": coverage_pct,
                                "gap_count": row.get("Gap Count", 0)
                            }
                    
                    # Store or update client data
                    if client_id not in gap_data:
                        gap_data[client_id] = {
                            "client_name": client_name,
                            "coverage_percentage": 0,
                            "gaps_count": len(gaps),
                            "total_rules": 0
                        }
                    
                    gap_data[client_id]["detailed_data"] = {
                        "gaps": gaps,
                        "tactic_coverage": tactic_coverage,
                        "file_path": str(gap_file)
                    }
                    
                    self.logger.info(f"Loaded detailed gap data for {client_name}: {len(gaps)} gaps")
                    
                except Exception as e:
                    self.logger.warning(f"Could not read detailed data from {gap_file}: {e}")
                    
            except Exception as e:
                self.logger.error(f"Failed to process gap file {gap_file}: {e}")
                
        return gap_data
        
    def analyze_common_gaps(self, gap_data: Dict[str, Dict]) -> Dict:
        """Analyze common gaps across all clients."""
        self.logger.info("Analyzing common gaps across all clients")
        
        # Count technique gaps across clients
        technique_gap_counts = Counter()
        technique_details = {}
        tactic_gap_analysis = defaultdict(lambda: {"total_clients": 0, "gaps_by_client": [], "avg_coverage": 0})
        client_coverage_stats = []
        
        total_clients = len([c for c in gap_data.values() if c.get("detailed_data")])
        
        for client_id, client_data in gap_data.items():
            detailed_data = client_data.get("detailed_data")
            if not detailed_data:
                continue
                
            client_name = client_data["client_name"]
            client_coverage_stats.append({
                "client_id": client_id,
                "client_name": client_name,
                "coverage_percentage": client_data.get("coverage_percentage", 0),
                "gaps_count": client_data.get("gaps_count", 0),
                "total_rules": client_data.get("total_rules", 0)
            })
            
            # Count technique gaps
            for gap in detailed_data["gaps"]:
                technique_id = gap["technique_id"]
                technique_gap_counts[technique_id] += 1
                
                # Store technique details
                if technique_id not in technique_details:
                    technique_details[technique_id] = {
                        "name": gap["technique_name"],
                        "tactics": gap["tactics"],
                        "priority": gap["priority"],
                        "affected_clients": []
                    }
                technique_details[technique_id]["affected_clients"].append(client_name)
            
            # Analyze tactic coverage
            for tactic_name, tactic_data in detailed_data["tactic_coverage"].items():
                tactic_gap_analysis[tactic_name]["total_clients"] += 1
                tactic_gap_analysis[tactic_name]["gaps_by_client"].append({
                    "client": client_name,
                    "coverage": tactic_data["coverage_percentage"],
                    "gaps": tactic_data["gap_count"]
                })
        
        # Calculate tactic averages
        for tactic_name, tactic_data in tactic_gap_analysis.items():
            if tactic_data["gaps_by_client"]:
                tactic_data["avg_coverage"] = sum(
                    client["coverage"] for client in tactic_data["gaps_by_client"]
                ) / len(tactic_data["gaps_by_client"])
                
        # Identify most common gaps
        most_common_gaps = []
        for technique_id, count in technique_gap_counts.most_common():
            if technique_id in technique_details:
                gap_percentage = (count / total_clients) * 100
                priority = self.determine_gap_priority(count, total_clients, technique_details[technique_id])
                
                most_common_gaps.append({
                    "technique_id": technique_id,
                    "technique_name": technique_details[technique_id]["name"],
                    "tactics": technique_details[technique_id]["tactics"],
                    "affected_clients": count,
                    "total_clients": total_clients,
                    "gap_percentage": gap_percentage,
                    "priority": priority,
                    "affected_client_names": technique_details[technique_id]["affected_clients"]
                })
        
        # Calculate overall statistics
        if client_coverage_stats:
            avg_coverage = sum(c["coverage_percentage"] for c in client_coverage_stats) / len(client_coverage_stats)
            total_gaps = sum(c["gaps_count"] for c in client_coverage_stats)
            total_rules = sum(c["total_rules"] for c in client_coverage_stats)
        else:
            avg_coverage = 0
            total_gaps = 0
            total_rules = 0
            
        return {
            "analysis_metadata": {
                "timestamp": datetime.now().isoformat(),
                "total_clients_analyzed": total_clients,
                "average_coverage": avg_coverage,
                "total_gaps_identified": total_gaps,
                "total_rules_analyzed": total_rules
            },
            "most_common_gaps": most_common_gaps,
            "tactic_analysis": dict(tactic_gap_analysis),
            "client_coverage_stats": client_coverage_stats,
            "technique_gap_counts": dict(technique_gap_counts)
        }
        
    def determine_gap_priority(self, affected_clients: int, total_clients: int, technique_info: Dict) -> str:
        """Determine priority level for a gap based on various factors."""
        gap_percentage = (affected_clients / total_clients) * 100
        tactics = technique_info.get("tactics", "").lower()
        
        # Critical: High impact tactics affecting most clients
        if gap_percentage >= 80 and any(tactic in tactics for tactic in ["initial-access", "execution", "persistence"]):
            return "Critical"
        
        # High: Either high percentage or critical tactics
        elif gap_percentage >= 60 or any(tactic in tactics for tactic in ["initial-access", "execution", "persistence", "privilege-escalation"]):
            return "High"
        
        # Medium: Moderate percentage or important tactics
        elif gap_percentage >= 40 or any(tactic in tactics for tactic in ["defense-evasion", "credential-access", "lateral-movement"]):
            return "Medium"
        
        # Low: Everything else
        else:
            return "Low"
            
    def create_common_gaps_report(self, analysis_data: Dict, gap_data: Dict, output_path: Path) -> bool:
        """Create comprehensive common gaps analysis report."""
        try:
            wb = Workbook()
            
            # Sheet 1: Executive Summary
            self.create_executive_summary_sheet(wb, analysis_data)
            
            # Sheet 2: Common Gaps Analysis
            self.create_common_gaps_sheet(wb, analysis_data)
            
            # Sheet 3: Tactic Analysis
            self.create_tactic_analysis_sheet(wb, analysis_data)
            
            # Sheet 4: Client Comparison
            self.create_client_comparison_sheet(wb, analysis_data)
            
            # Sheet 5: Priority Recommendations
            self.create_priority_recommendations_sheet(wb, analysis_data)
            
            # Sheet 6: Implementation Roadmap
            self.create_implementation_roadmap_sheet(wb, analysis_data)
            
            # Save workbook
            wb.save(output_path)
            
            # Validate file creation
            if output_path.exists() and output_path.stat().st_size > 0:
                self.logger.info(f"Common gaps report created: {output_path}")
                return True
            else:
                self.logger.error(f"Failed to create common gaps report: {output_path}")
                return False
                
        except Exception as e:
            self.logger.error(f"Failed to create common gaps report: {e}")
            return False
            
    def create_executive_summary_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create executive summary sheet."""
        ws = wb.active
        ws.title = "Executive Summary"
        
        metadata = analysis_data["analysis_metadata"]
        most_common_gaps = analysis_data["most_common_gaps"]
        
        # Executive summary data
        summary_data = [
            ["S-Rank Core - Common Gaps Analysis Report", ""],
            ["", ""],
            ["Executive Summary", ""],
            ["Analysis Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Total Clients Analyzed", metadata["total_clients_analyzed"]],
            ["Average Coverage Across Clients", f"{metadata['average_coverage']:.1f}%"],
            ["Total Detection Rules Analyzed", metadata["total_rules_analyzed"]],
            ["Total Gaps Identified", metadata["total_gaps_identified"]],
            ["", ""],
            ["Key Findings", ""],
        ]
        
        # Add top critical gaps
        critical_gaps = [gap for gap in most_common_gaps if gap["priority"] == "Critical"][:5]
        if critical_gaps:
            summary_data.append(["Critical Gaps (Immediate Action Required):", ""])
            for i, gap in enumerate(critical_gaps, 1):
                summary_data.append([
                    f"  {i}. {gap['technique_id']} - {gap['technique_name']}",
                    f"Affects {gap['affected_clients']}/{gap['total_clients']} clients ({gap['gap_percentage']:.0f}%)"
                ])
        
        # Add high priority gaps
        high_gaps = [gap for gap in most_common_gaps if gap["priority"] == "High"][:5]
        if high_gaps:
            summary_data.extend([
                ["", ""],
                ["High Priority Gaps:", ""]
            ])
            for i, gap in enumerate(high_gaps, 1):
                summary_data.append([
                    f"  {i}. {gap['technique_id']} - {gap['technique_name']}",
                    f"Affects {gap['affected_clients']}/{gap['total_clients']} clients ({gap['gap_percentage']:.0f}%)"
                ])
        
        # Add recommendations summary
        summary_data.extend([
            ["", ""],
            ["Strategic Recommendations", ""],
            ["1. Immediate Focus Areas:", ""],
            ["   • Address Critical gaps affecting 80%+ of clients", ""],
            ["   • Prioritize Initial Access and Execution techniques", ""],
            ["   • Implement cross-client detection sharing", ""],
            ["", ""],
            ["2. Medium-term Initiatives:", ""],
            ["   • Develop standardized detection templates", ""],
            ["   • Establish gap monitoring and tracking", ""],
            ["   • Create client-specific improvement plans", ""],
            ["", ""],
            ["3. Long-term Strategy:", ""],
            ["   • Build automated gap detection and remediation", ""],
            ["   • Establish continuous coverage monitoring", ""],
            ["   • Develop threat-informed detection priorities", ""]
        ])
        
        # Write data to sheet
        for row_idx, row_data in enumerate(summary_data, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Style different types of content
                if col_idx == 1 and value and not str(value).startswith("  "):
                    if row_idx == 1:  # Main title
                        cell.font = Font(bold=True, size=18, color="702386")
                    elif "Summary" in str(value) or "Findings" in str(value) or "Recommendations" in str(value):
                        cell.font = Font(bold=True, size=14)
                    elif any(char.isdigit() for char in str(value)[:3]):  # Numbered items
                        cell.font = Font(bold=True)
                elif "Critical" in str(value):
                    cell.fill = self.critical_fill
                elif "High Priority" in str(value):
                    cell.fill = self.high_fill
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 40
        
    def create_common_gaps_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create common gaps analysis sheet."""
        ws = wb.create_sheet(title="Common Gaps Analysis")
        
        # Headers
        headers = [
            "Technique ID", "Technique Name", "Tactics", "Affected Clients", 
            "Total Clients", "Gap %", "Priority", "Affected Client Names"
        ]
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Add common gaps data
        most_common_gaps = analysis_data["most_common_gaps"]
        
        for row_idx, gap in enumerate(most_common_gaps, start=2):
            ws.cell(row=row_idx, column=1, value=gap["technique_id"])
            ws.cell(row=row_idx, column=2, value=gap["technique_name"])
            ws.cell(row=row_idx, column=3, value=gap["tactics"])
            ws.cell(row=row_idx, column=4, value=gap["affected_clients"])
            ws.cell(row=row_idx, column=5, value=gap["total_clients"])
            
            # Gap percentage with color coding
            gap_pct_cell = ws.cell(row=row_idx, column=6, value=f"{gap['gap_percentage']:.1f}%")
            
            # Priority with color coding
            priority_cell = ws.cell(row=row_idx, column=7, value=gap["priority"])
            
            # Color code based on priority
            if gap["priority"] == "Critical":
                gap_pct_cell.fill = self.critical_fill
                priority_cell.fill = self.critical_fill
            elif gap["priority"] == "High":
                gap_pct_cell.fill = self.high_fill
                priority_cell.fill = self.high_fill
            elif gap["priority"] == "Medium":
                gap_pct_cell.fill = self.medium_fill
                priority_cell.fill = self.medium_fill
            else:
                gap_pct_cell.fill = self.low_fill
                priority_cell.fill = self.low_fill
            
            # Affected clients list
            ws.cell(row=row_idx, column=8, value=", ".join(gap["affected_client_names"]))
        
        # Auto-adjust column widths
        widths = [15, 35, 25, 15, 15, 10, 12, 40]
        for col_idx, width in enumerate(widths, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
            
        # Add auto-filter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(most_common_gaps) + 1}"
        
    def create_tactic_analysis_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create tactic analysis sheet."""
        ws = wb.create_sheet(title="Tactic Analysis")
        
        # Headers
        headers = ["Tactic", "Avg Coverage %", "Clients Analyzed", "Lowest Coverage", "Highest Coverage", "Std Deviation"]
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Analyze tactic data
        tactic_analysis = analysis_data["tactic_analysis"]
        tactic_summary = []
        
        for tactic_name, tactic_data in tactic_analysis.items():
            if tactic_data["gaps_by_client"]:
                coverages = [client["coverage"] for client in tactic_data["gaps_by_client"]]
                
                tactic_summary.append({
                    "tactic": tactic_name,
                    "avg_coverage": tactic_data["avg_coverage"],
                    "clients_count": len(coverages),
                    "min_coverage": min(coverages),
                    "max_coverage": max(coverages),
                    "std_dev": (sum((x - tactic_data["avg_coverage"]) ** 2 for x in coverages) / len(coverages)) ** 0.5 if len(coverages) > 1 else 0
                })
        
        # Sort by average coverage (lowest first)
        tactic_summary.sort(key=lambda x: x["avg_coverage"])
        
        # Add tactic data
        for row_idx, tactic in enumerate(tactic_summary, start=2):
            ws.cell(row=row_idx, column=1, value=tactic["tactic"])
            
            avg_coverage_cell = ws.cell(row=row_idx, column=2, value=f"{tactic['avg_coverage']:.1f}%")
            
            # Color code based on average coverage
            if tactic["avg_coverage"] < 30:
                avg_coverage_cell.fill = self.critical_fill
            elif tactic["avg_coverage"] < 50:
                avg_coverage_cell.fill = self.high_fill
            elif tactic["avg_coverage"] < 70:
                avg_coverage_cell.fill = self.medium_fill
            else:
                avg_coverage_cell.fill = self.low_fill
            
            ws.cell(row=row_idx, column=3, value=tactic["clients_count"])
            ws.cell(row=row_idx, column=4, value=f"{tactic['min_coverage']:.1f}%")
            ws.cell(row=row_idx, column=5, value=f"{tactic['max_coverage']:.1f}%")
            ws.cell(row=row_idx, column=6, value=f"{tactic['std_dev']:.1f}")
        
        # Auto-adjust column widths
        widths = [25, 18, 18, 18, 18, 15]
        for col_idx, width in enumerate(widths, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
            
    def create_client_comparison_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create client comparison sheet."""
        ws = wb.create_sheet(title="Client Comparison")
        
        # Headers
        headers = ["Client Name", "Coverage %", "Total Gaps", "Total Rules", "Relative Performance", "Focus Areas"]
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        client_stats = analysis_data["client_coverage_stats"]
        avg_coverage = analysis_data["analysis_metadata"]["average_coverage"]
        
        # Sort clients by coverage percentage
        sorted_clients = sorted(client_stats, key=lambda x: x["coverage_percentage"], reverse=True)
        
        for row_idx, client in enumerate(sorted_clients, start=2):
            ws.cell(row=row_idx, column=1, value=client["client_name"])
            
            coverage_cell = ws.cell(row=row_idx, column=2, value=f"{client['coverage_percentage']:.1f}%")
            
            # Color code coverage
            if client["coverage_percentage"] >= avg_coverage + 10:
                coverage_cell.fill = self.low_fill
                performance = "Above Average"
            elif client["coverage_percentage"] >= avg_coverage - 10:
                coverage_cell.fill = self.medium_fill
                performance = "Average"
            elif client["coverage_percentage"] >= avg_coverage - 20:
                coverage_cell.fill = self.high_fill
                performance = "Below Average"
            else:
                coverage_cell.fill = self.critical_fill
                performance = "Needs Attention"
            
            ws.cell(row=row_idx, column=3, value=client["gaps_count"])
            ws.cell(row=row_idx, column=4, value=client["total_rules"])
            ws.cell(row=row_idx, column=5, value=performance)
            
            # Determine focus areas based on performance
            if performance == "Needs Attention":
                focus_areas = "Critical gaps, Basic detection coverage"
            elif performance == "Below Average":
                focus_areas = "High-priority gaps, Technique expansion"
            elif performance == "Average":
                focus_areas = "Medium-priority gaps, Coverage optimization"
            else:
                focus_areas = "Advanced techniques, False positive reduction"
                
            ws.cell(row=row_idx, column=6, value=focus_areas)
        
        # Auto-adjust column widths
        widths = [25, 15, 15, 15, 20, 35]
        for col_idx, width in enumerate(widths, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
            
    def create_priority_recommendations_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create priority recommendations sheet."""
        ws = wb.create_sheet(title="Priority Recommendations")
        
        most_common_gaps = analysis_data["most_common_gaps"]
        
        recommendations = [
            ["Priority-Based Detection Recommendations", ""],
            ["", ""],
            ["CRITICAL PRIORITY (Immediate Action - 0-30 days)", ""],
        ]
        
        # Critical recommendations
        critical_gaps = [gap for gap in most_common_gaps if gap["priority"] == "Critical"]
        if critical_gaps:
            for i, gap in enumerate(critical_gaps[:10], 1):
                recommendations.append([
                    f"{i}. {gap['technique_id']} - {gap['technique_name']}",
                    f"Affects {gap['affected_clients']}/{gap['total_clients']} clients"
                ])
                recommendations.append([
                    f"   Tactics: {gap['tactics']}", 
                    f"Priority: {gap['priority']}"
                ])
                recommendations.append([
                    f"   Recommended Action: Implement centralized detection rule",
                    f"Expected Impact: High"
                ])
                recommendations.append(["", ""])
        
        # High priority recommendations
        recommendations.extend([
            ["HIGH PRIORITY (30-90 days)", ""],
        ])
        
        high_gaps = [gap for gap in most_common_gaps if gap["priority"] == "High"]
        if high_gaps:
            for i, gap in enumerate(high_gaps[:8], 1):
                recommendations.append([
                    f"{i}. {gap['technique_id']} - {gap['technique_name']}",
                    f"Affects {gap['affected_clients']}/{gap['total_clients']} clients"
                ])
                
        # Medium priority recommendations
        recommendations.extend([
            ["", ""],
            ["MEDIUM PRIORITY (90-180 days)", ""],
        ])
        
        medium_gaps = [gap for gap in most_common_gaps if gap["priority"] == "Medium"]
        if medium_gaps:
            for i, gap in enumerate(medium_gaps[:5], 1):
                recommendations.append([
                    f"{i}. {gap['technique_id']} - {gap['technique_name']}",
                    f"Affects {gap['affected_clients']}/{gap['total_clients']} clients"
                ])
        
        # Implementation guidance
        recommendations.extend([
            ["", ""],
            ["IMPLEMENTATION GUIDANCE", ""],
            ["", ""],
            ["Detection Development Approach:", ""],
            ["1. Critical Gaps:", ""],
            ["   • Develop high-fidelity rules with low false positives", ""],
            ["   • Test across multiple client environments", ""],
            ["   • Implement with immediate alerting", ""],
            ["", ""],
            ["2. High Priority Gaps:", ""],
            ["   • Focus on behavioral detection methods", ""],
            ["   • Consider environmental context", ""],
            ["   • Plan for tuning and optimization", ""],
            ["", ""],
            ["3. Medium Priority Gaps:", ""],
            ["   • Leverage threat intelligence integration", ""],
            ["   • Consider hunt-based detection approaches", ""],
            ["   • Plan for gradual rollout", ""]
        ])
        
        # Write recommendations to sheet
        for row_idx, row_data in enumerate(recommendations, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Style based on content
                if col_idx == 1 and value:
                    if "CRITICAL" in str(value):
                        cell.fill = self.critical_fill
                        cell.font = Font(bold=True, color="FFFFFF")
                    elif "HIGH PRIORITY" in str(value):
                        cell.fill = self.high_fill
                        cell.font = Font(bold=True, color="000000")
                    elif "MEDIUM PRIORITY" in str(value):
                        cell.fill = self.medium_fill
                        cell.font = Font(bold=True, color="000000")
                    elif "IMPLEMENTATION" in str(value):
                        cell.font = Font(bold=True, size=14)
                    elif any(char.isdigit() for char in str(value)[:3]):
                        cell.font = Font(bold=True)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 55
        ws.column_dimensions['B'].width = 30
        
    def create_implementation_roadmap_sheet(self, wb: Workbook, analysis_data: Dict):
        """Create implementation roadmap sheet."""
        ws = wb.create_sheet(title="Implementation Roadmap")
        
        roadmap_data = [
            ["S-Rank Core - Detection Implementation Roadmap", ""],
            ["", ""],
            ["PHASE 1: Critical Gaps (0-30 days)", ""],
            ["Goal: Address gaps affecting 80%+ of clients", ""],
            ["", ""],
            ["Week 1-2: Assessment and Planning", ""],
            ["• Validate critical gap findings with security teams", ""],
            ["• Prioritize techniques based on threat landscape", ""],
            ["• Establish detection development resources", ""],
            ["• Create testing and validation framework", ""],
            ["", ""],
            ["Week 3-4: Detection Development", ""],
            ["• Develop high-fidelity detection rules", ""],
            ["• Test rules in lab environment", ""],
            ["• Document rule logic and requirements", ""],
            ["• Create deployment packages", ""],
            ["", ""],
            ["PHASE 2: High Priority Gaps (30-90 days)", ""],
            ["Goal: Improve coverage for high-impact techniques", ""],
            ["", ""],
            ["Month 2: Behavioral Detection", ""],
            ["• Implement behavioral analytics rules", ""],
            ["• Deploy machine learning models", ""],
            ["• Establish baseline behavioral patterns", ""],
            ["• Create correlation rules", ""],
            ["", ""],
            ["Month 3: Environmental Tuning", ""],
            ["• Customize rules for client environments", ""],
            ["• Optimize for false positive reduction", ""],
            ["• Implement advanced hunting queries", ""],
            ["• Establish monitoring dashboards", ""],
            ["", ""],
            ["PHASE 3: Medium Priority & Optimization (90-180 days)", ""],
            ["Goal: Comprehensive coverage and continuous improvement", ""],
            ["", ""],
            ["Month 4-5: Coverage Expansion", ""],
            ["• Address remaining medium-priority gaps", ""],
            ["• Implement threat intelligence integration", ""],
            ["• Develop custom detection logic", ""],
            ["• Create automated response workflows", ""],
            ["", ""],
            ["Month 6: Continuous Monitoring", ""],
            ["• Establish automated gap analysis", ""],
            ["• Implement coverage monitoring dashboards", ""],
            ["• Create performance metrics and KPIs", ""],
            ["• Develop improvement feedback loops", ""],
            ["", ""],
            ["SUCCESS METRICS", ""],
            ["", ""],
            ["Phase 1 Success Criteria:", ""],
            ["• 90%+ coverage for critical techniques", ""],
            ["• <5% false positive rate", ""],
            ["• All critical gaps addressed", ""],
            ["", ""],
            ["Phase 2 Success Criteria:", ""],
            ["• 80%+ coverage for high-priority techniques", ""],
            ["• Advanced behavioral detection deployed", ""],
            ["• Client-specific optimizations complete", ""],
            ["", ""],
            ["Phase 3 Success Criteria:", ""],
            ["• 70%+ overall MITRE coverage", ""],
            ["• Automated monitoring and alerting", ""],
            ["• Continuous improvement process established", ""]
        ]
        
        # Write roadmap data
        for row_idx, row_data in enumerate(roadmap_data, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                
                # Style based on content
                if col_idx == 1 and value:
                    if "PHASE" in str(value):
                        cell.fill = self.header_fill
                        cell.font = Font(bold=True, color="FFFFFF")
                    elif "SUCCESS METRICS" in str(value):
                        cell.font = Font(bold=True, size=14)
                    elif "Week" in str(value) or "Month" in str(value) or "Goal:" in str(value):
                        cell.font = Font(bold=True)
                    elif str(value).startswith("•"):
                        cell.font = Font(color="444444")
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 60
        ws.column_dimensions['B'].width = 20
        
    def scan_common_gaps(self) -> bool:
        """Main method to scan and analyze common gaps across all clients."""
        print(f"{Color.BOLD}🚀 S-Rank Core - Stage 4: Enhanced Common Gaps Scanner{Color.ENDC}")
        print(f"{Color.CYAN}📋 Configuration Summary:{Color.ENDC}")
        self.config_manager.print_config_summary()
        print(f"\n{Color.CYAN}📂 Input folder:{Color.ENDC} {self.latest_reports_folder}")
        print(f"{Color.CYAN}📂 Output folder:{Color.ENDC} {self.common_gaps_dir}")
        print()
        
        # Load gap analysis data from all clients
        print(f"{Color.CYAN}📥 Loading gap analysis data from all clients...{Color.ENDC}")
        gap_data = self.load_gap_analysis_data()
        
        if not gap_data:
            print(f"{Color.FAIL}❌ No gap analysis data found. Please ensure Stage 3 (Gap Analysis) has been completed.{Color.ENDC}")
            return False
            
        clients_with_data = len([c for c in gap_data.values() if c.get("detailed_data")])
        print(f"{Color.OKGREEN}✅ Loaded gap data for {clients_with_data} clients{Color.ENDC}")
        
        if clients_with_data == 0:
            print(f"{Color.FAIL}❌ No detailed gap data available for analysis.{Color.ENDC}")
            return False
        
        # Analyze common gaps
        print(f"{Color.CYAN}🔍 Analyzing common gaps across all clients...{Color.ENDC}")
        analysis_data = self.analyze_common_gaps(gap_data)
        
        # Generate output filename
        output_filename = f"common_gaps_report_{self.pipeline_timestamp}.xlsx"
        output_path = self.common_gaps_dir / output_filename
        
        # Create comprehensive report
        print(f"{Color.CYAN}📊 Creating comprehensive common gaps report...{Color.ENDC}")
        if self.create_common_gaps_report(analysis_data, gap_data, output_path):
            
            # Create summary JSON
            summary_path = self.common_gaps_dir / "common_gaps_summary_enhanced.json"
            self.create_summary_json(analysis_data, summary_path)
            
            # Display results
            metadata = analysis_data["analysis_metadata"]
            most_common_gaps = analysis_data["most_common_gaps"]
            
            print(f"\n{Color.BOLD}📊 Common Gaps Analysis Results:{Color.ENDC}")
            print(f"   • Clients analyzed: {metadata['total_clients_analyzed']}")
            print(f"   • Average coverage: {metadata['average_coverage']:.1f}%")
            print(f"   • Total gaps identified: {metadata['total_gaps_identified']}")
            print(f"   • Unique techniques with gaps: {len(most_common_gaps)}")
            
            # Show top gaps
            critical_gaps = [gap for gap in most_common_gaps if gap["priority"] == "Critical"]
            high_gaps = [gap for gap in most_common_gaps if gap["priority"] == "High"]
            
            if critical_gaps:
                print(f"\n{Color.FAIL}🔴 Critical Gaps (affecting 80%+ clients):{Color.ENDC}")
                for gap in critical_gaps[:5]:
                    print(f"   • {gap['technique_id']}: {gap['gap_percentage']:.0f}% of clients affected")
                    
            if high_gaps:
                print(f"\n{Color.WARNING}🟡 High Priority Gaps:{Color.ENDC}")
                for gap in high_gaps[:5]:
                    print(f"   • {gap['technique_id']}: {gap['gap_percentage']:.0f}% of clients affected")
            
            print(f"\n{Color.OKGREEN}✅ Common gaps report created: {output_filename}{Color.ENDC}")
            print(f"{Color.CYAN}📄 Summary saved to: common_gaps_summary_enhanced.json{Color.ENDC}")
            
            return True
        else:
            print(f"{Color.FAIL}❌ Failed to create common gaps report{Color.ENDC}")
            return False
            
    def create_summary_json(self, analysis_data: Dict, summary_path: Path):
        """Create JSON summary of common gaps analysis."""
        try:
            summary_data = {
                "common_gaps_summary": {
                    "timestamp": datetime.now().isoformat(),
                    "pipeline_timestamp": self.pipeline_timestamp,
                    "processing_version": "2.0_enhanced_common_gaps",
                    **analysis_data["analysis_metadata"]
                },
                "critical_gaps": [
                    gap for gap in analysis_data["most_common_gaps"] 
                    if gap["priority"] == "Critical"
                ],
                "high_priority_gaps": [
                    gap for gap in analysis_data["most_common_gaps"] 
                    if gap["priority"] == "High"
                ],
                "client_comparison": analysis_data["client_coverage_stats"],
                "configuration": {
                    "technique_naming_format": self.config_manager.get_technique_naming_format(),
                    "processing_settings": self.processing_settings
                }
            }
            
            with open(summary_path, 'w', encoding='utf-8') as f:
                json.dump(summary_data, f, indent=2, ensure_ascii=False)
                
            self.logger.info(f"Common gaps summary saved to: {summary_path}")
            
        except Exception as e:
            self.logger.error(f"Failed to create summary JSON: {e}")


def main():
    """Entry point for the enhanced common gaps scanner script."""
    try:
        scanner = EnhancedGapsScanner()
        success = scanner.scan_common_gaps()
        return 0 if success else 1
    except KeyboardInterrupt:
        print(f"\n{Color.WARNING}⚠️  Operation cancelled by user.{Color.ENDC}")
        return 130
    except Exception as e:
        print(f"{Color.FAIL}💥 Fatal error: {e}{Color.ENDC}")
        logging.error(f"Fatal error: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())